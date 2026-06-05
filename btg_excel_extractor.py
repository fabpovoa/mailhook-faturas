"""
Extrai transações de uma fatura BTG Pactual no formato Excel (.xlsx protegido).

Estrutura do arquivo (aba "Titular"):
  - Múltiplas seções de transações, cada uma precedida por uma linha de cabeçalho
    contendo "Data", "Descrição", "Valor" e "Final Cartão".
  - As colunas variam entre seções (ex: seção de créditos não tem "Tipo de compra").
  - O extrator detecta TODOS os cabeçalhos e processa cada seção.

Mapeamento Final Cartão → aba:
  0627 = Luiz
  7387 = Fátima 7387
  5310 = Fátima 5310
  1205 = Hiero
"""
import io
import re
import logging
from datetime import datetime
from typing import Optional

import msoffcrypto
import openpyxl

log = logging.getLogger(__name__)

CARD_TO_TAB = {
    "0627": "Luiz",
    "7387": "Fátima 7387",
    "5310": "Fátima 5310",
    "1205": "Hiero",
}

PARCELA_RE = re.compile(r"\((\d+/\d+)\)")

MES_MAP = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8,
    "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}


def _unlock(xlsx_bytes: bytes, password: str) -> bytes:
    try:
        enc = msoffcrypto.OfficeFile(io.BytesIO(xlsx_bytes))
        enc.load_key(password=password)
        buf = io.BytesIO()
        enc.decrypt(buf)
        return buf.getvalue()
    except msoffcrypto.exceptions.DecryptionError:
        # Arquivo já desbloqueado — retorna como está
        return xlsx_bytes


def _parse_ref_date(ws) -> Optional[datetime]:
    for row in ws.iter_rows(max_row=10, values_only=True):
        for cell in row:
            if cell and isinstance(cell, str):
                m = re.search(r"(\w+)/(\d{4})", cell)
                if m:
                    mes = MES_MAP.get(m.group(1).lower())
                    if mes:
                        return datetime(int(m.group(2)), mes, 1)
    return None


def _map_header(row_vals: list) -> Optional[dict]:
    """
    Dado os valores de uma linha, retorna dict com índices 1-based das colunas
    se for uma linha de cabeçalho de transações (tem Data + Final Cartão).
    Retorna None se não for cabeçalho.
    """
    lower = [str(v or "").lower() for v in row_vals]
    if "data" not in lower or not any("final" in v for v in lower):
        return None

    col = {}
    for i, v in enumerate(lower, 1):
        if v == "data" and "data" not in col:
            col["data"] = i
        elif "descri" in v and "desc" not in col:
            col["desc"] = i
        elif v == "valor" and "valor" not in col:
            col["valor"] = i
        elif "tipo" in v and "tipo" not in col:
            col["tipo"] = i
        elif ("código" in v or "codigo" in v or "autoriza" in v) and "cod" not in col:
            col["cod"] = i
        elif "final" in v and "final" not in col:
            col["final"] = i

    if "data" in col and "valor" in col and "final" in col:
        return col
    return None


def extract_btg_excel(xlsx_bytes: bytes, password: str) -> tuple[dict[str, list[dict]], Optional[datetime]]:
    unlocked = _unlock(xlsx_bytes, password)
    wb = openpyxl.load_workbook(io.BytesIO(unlocked))
    ws = wb.active

    ref_date = _parse_ref_date(ws)
    grouped: dict[str, list[dict]] = {tab: [] for tab in CARD_TO_TAB.values()}
    seen: set[tuple] = set()  # deduplicação

    all_rows = list(ws.iter_rows(values_only=True))
    current_col_map = None

    for row_vals in all_rows:
        # Verifica se é linha de cabeçalho
        col_map = _map_header(list(row_vals))
        if col_map:
            current_col_map = col_map
            log.debug("Novo cabeçalho detectado: %s", col_map)
            continue

        if not current_col_map:
            continue

        def get(key):
            idx = current_col_map.get(key)
            return row_vals[idx - 1] if idx and idx <= len(row_vals) else None

        data_val  = get("data")
        desc_val  = get("desc")
        valor_val = get("valor")
        final_val = str(get("final") or "").strip()

        # Só processa linhas com data datetime e Final Cartão conhecido
        if not isinstance(data_val, datetime):
            continue
        if not desc_val or valor_val is None:
            continue

        tab = CARD_TO_TAB.get(final_val)
        if not tab:
            continue

        data_str  = data_val.strftime("%d/%m/%Y")
        desc_str  = str(desc_val).strip()
        valor_flt = float(valor_val)

        # Deduplicação por (data, desc, valor, cartão)
        key = (data_str, desc_str, valor_flt, final_val)
        if key in seen:
            continue
        seen.add(key)

        parcela_m = PARCELA_RE.search(desc_str)
        parcela   = parcela_m.group(1) if parcela_m else ""
        tipo      = str(get("tipo") or "Compra").strip()
        cod       = str(get("cod")  or "").strip()

        grouped[tab].append({
            "Data":             data_str,
            "Descrição":        desc_str,
            "Parcela":          parcela,
            "Valor (R$)":       valor_flt,
            "Tipo":             tipo,
            "Cod. Autorização": cod,
            "Final Cartão":     final_val,
        })

    totals = {k: len(v) for k, v in grouped.items()}
    log.info("BTG Excel — extração concluída: %s (ref: %s)", totals, ref_date)
    return grouped, ref_date
