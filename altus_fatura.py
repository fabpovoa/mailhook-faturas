#!/usr/bin/env python3
"""Fatura BB Altus Visa — automação mensal, sem LLM (launchd roda todo dia; cada etapa é idempotente).

E-mail (Gmail) -> PDF desbloqueado em Cartões/{ANO}/{MM}/BB_Altus -> xlsx dos pais (abas Luiz/Mariana/Fátima)
-> Boletos (gsheet) -> Contas a Pagar (money) -> planilha de pagtos -> e-mail à família -> WhatsApp (grupo).

Uso: altus_fatura.py [--month YYYY-MM] [--dry-run] [--mark-done]
Estado por mês em altus_state.json (etapa concluída nunca repete). Spec: ~/.claude/skills/altus/SKILL.md
"""
import argparse, base64, json, os, re, subprocess, sys, tempfile, unicodedata, urllib.request
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

os.environ["PATH"] = "/usr/local/bin:/opt/homebrew/bin:" + os.environ.get("PATH", "")  # launchd começa sem PATH
HERE = Path(__file__).resolve().parent
os.chdir(HERE)  # gmail.py lê token.json relativo
sys.path.insert(0, str(HERE))

from dotenv import dotenv_values, load_dotenv
load_dotenv(HERE / ".env")

BASE = Path("/Users/fabpovoa/Library/CloudStorage/GoogleDrive-fabio@smartmoney.ventures/My Drive/Personal stuff")
CARTOES, TEMPLATES = BASE / "Cartões", BASE / "Smart Money Kpool Ventures/Templates"
KPOOL = Path("/Users/fabpovoa/Developer/kpool-app")
MONEY = Path("/Users/fabpovoa/Developer/money")
STATE = HERE / "altus_state.json"
BOLETOS_ID, BOLETOS_TAB = "1G3S57VJmz-3EQIwK5Lo_iuJhVaI7X2h1EtPvPBpLxZo", "Acomp mensal May/26 em diante"
PASTAS = ["01_jan", "02_fev", "03_mar", "04_abr", "05_may", "06_jun", "07_jul", "08_ago", "09_set", "10_out", "11_nov", "12_dez"]
MES = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
ABAS = {"Luiz": ("BB Altus Visa Luiz", "LUIZ"), "Mariana": ("BB Altus Visa Mariana", "MARIANA"), "Fátima": ("BB Altus Visa Fátima", "MARIA F")}
SENDER, SUBJECT = "digital@faturaourocard.com.br", "FATURA CLIENTE - CARTAO ALTUS VISA"


class Pendente(Exception):
    """Ainda não dá pra fazer (e-mail não chegou, Drive não sincronizou); tenta de novo na próxima rodada."""


def log(*a):
    print(datetime.now().strftime("%Y-%m-%d %H:%M:%S"), *a, flush=True)


def notify(msg):
    log("NOTIFY", msg)
    subprocess.run(["osascript", "-e", f'display notification "{msg[:180]}" with title "Fatura Altus"'], check=False)


def brl(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def dia_util(d):
    """Primeiro dia útil >= d (fim de semana ou feriado nacional empurra pra frente)."""
    # Páscoa (algoritmo gregoriano anônimo)
    y = d.year; a_ = y % 19; b_, c_ = divmod(y, 100); d_, e_ = divmod(b_, 4); f_ = (b_ + 8) // 25
    g_ = (b_ - f_ + 1) // 3; h_ = (19 * a_ + b_ - d_ - g_ + 15) % 30; i_, k_ = divmod(c_, 4)
    l_ = (32 + 2 * e_ + 2 * i_ - h_ - k_) % 7; m_ = (a_ + 11 * h_ + 22 * l_) // 451
    mes, dia = divmod(h_ + l_ - 7 * m_ + 114, 31)
    pascoa = date(y, mes, dia + 1)
    feriados = {date(y, mm, dd) for mm, dd in [(1, 1), (4, 21), (5, 1), (9, 7), (10, 12), (11, 2), (11, 15), (11, 20), (12, 25)]}
    feriados |= {pascoa - timedelta(48), pascoa - timedelta(47), pascoa - timedelta(2), pascoa + timedelta(60)}
    while d.weekday() >= 5 or d in feriados:
        d += timedelta(1)
    return d


# ---------------------------------------------------------------- parser do PDF
def parse_pdf(pdf):
    txt = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
    return txt, parse_secoes(txt)


def parse_secoes(txt, ref_month=None, ref_year=None):
    venc = re.search(r"(\d\d)/(\d\d)/(\d{4})", txt[txt.index("Vencimento"):])
    ref_month, ref_year = ref_month or int(venc[2]), ref_year or int(venc[3])
    secs, cur, cat = {}, None, None
    sec_re = re.compile(r"^\s*(\d\d)- (.+?)\s+Cartao N\. (\d+)")
    tx_re = re.compile(r"^\s*(\d\d)/(\d\d)\s+(.+?)\s+([A-Z]{2})\s+R\$\s*([\d.]+,\d\d)(-?)\s*$")
    sub_re = re.compile(r"^\s*Subtotal\s+R\$\s*([\d.]+,\d\d)(-?)")
    for line in txt.splitlines():
        if not line.strip() or line.startswith("Página") or re.match(r"\s*Data\s+Descri", line):
            continue
        if m := sec_re.match(line):
            cur = secs.setdefault(m[2].strip(), {"tx": [], "subtotal": None}); cat = None; continue
        if cur is None:
            continue
        if m := sub_re.match(line):
            cur["subtotal"] = Decimal(m[1].replace(".", "").replace(",", ".")) * (-1 if m[2] else 1); continue
        if m := tx_re.match(line):
            d, mo, desc, pais, val, neg = m.groups()
            ano = ref_year - 1 if int(mo) > ref_month else ref_year
            v = Decimal(val.replace(".", "").replace(",", ".")) * (-1 if neg else 1)
            p = re.split(r"\s{2,}", desc.strip(), maxsplit=1)
            cur["tx"].append(dict(data=f"{d}.{mo}.{ano}", desc=p[0], cidade=p[1] if len(p) > 1 else None, pais=pais, valor=v, cat=cat))
            continue
        if not re.search(r"US\$|\*\*\*|equivalente|Cotação|^\s*Total", line):
            cat = line.strip()
    return secs


def secao(secs, chave):
    ach = [v for k, v in secs.items() if chave in k]
    if len(ach) != 1:
        raise RuntimeError(f"seção '{chave}' não encontrada/ambígua no PDF ({list(secs)})")
    return ach[0]


# ---------------------------------------------------------------- Google
def creds():
    from gmail import get_credentials
    return get_credentials()


def drive_link(name):
    from googleapiclient.discovery import build
    r = build("drive", "v3", credentials=creds()).files().list(
        q=f"name = '{name}' and trashed = false", fields="files(webViewLink,modifiedTime)", orderBy="modifiedTime desc").execute()
    if not r["files"]:
        raise Pendente(f"'{name}' ainda não apareceu no Drive (sincronizando)")
    return r["files"][0]["webViewLink"]


def limpa_amarelo(sh, ws, celulas):
    """celulas: lista (linha1, coluna1) — remove só o fundo, mantém o resto da formatação."""
    sh.batch_update({"requests": [{"repeatCell": {"range": {"sheetId": ws.id, "startRowIndex": r - 1, "endRowIndex": r, "startColumnIndex": c - 1, "endColumnIndex": c},
                                                   "cell": {"userEnteredFormat": {}}, "fields": "userEnteredFormat.backgroundColor"}} for r, c in celulas]})


# ---------------------------------------------------------------- contexto
class Ctx:
    def __init__(self, ano, mes, dry):
        self.ano, self.mes, self.dry = ano, mes, dry
        self.ym = f"{ano}-{mes:02d}"
        self.pasta_mes = CARTOES / str(ano) / PASTAS[mes - 1]
        self.dest = self.pasta_mes / "BB_Altus" / f"CARTAO_ALTUS_VISA-{MES[mes - 1]}_{str(ano)[2:]}.pdf"
        allst = json.loads(STATE.read_text()) if STATE.exists() else {}
        self.st = allst.setdefault(self.ym, {"steps": {}, "d": {}})
        self._all = allst

    def save(self):
        if not self.dry:
            STATE.write_text(json.dumps(self._all, indent=1, ensure_ascii=False))

    @property
    def d(self):
        return self.st["d"]

    def xlsx(self):
        a = [p for p in (self.pasta_mes / "Pais").glob("faturas_Cartões_Luiz_*.xlsx") if not p.name.startswith("~$")]
        if len(a) != 1:
            raise RuntimeError(f"esperava 1 xlsx dos pais em {self.pasta_mes / 'Pais'}, achei {len(a)}")
        return a[0]


# ---------------------------------------------------------------- etapas
def etapa_pdf(c):
    if not c.dest.exists():
        from gmail import get_gmail_service, extract_pdf_attachments
        s = get_gmail_service()
        q = f'from:{SENDER} subject:"{SUBJECT}" has:attachment after:{c.ano}/{c.mes:02d}/01'
        ms = s.users().messages().list(userId="me", q=q, maxResults=5).execute().get("messages", [])
        if not ms:
            raise Pendente("e-mail da fatura ainda não chegou")
        msg = s.users().messages().get(userId="me", id=ms[0]["id"], format="full").execute()
        anexos = extract_pdf_attachments(s, "me", msg)
        if not anexos:
            raise RuntimeError("e-mail da fatura sem PDF anexo")
        with tempfile.TemporaryDirectory() as t:
            lock, unlocked = Path(t, "lock.pdf"), Path(t, "ok.pdf")
            lock.write_bytes(anexos[0]["data"])
            r = subprocess.run(["qpdf", f"--password={os.environ['ALTUS_PDF_PASSWORD']}", "--decrypt", str(lock), str(unlocked)], capture_output=True, text=True)
            if r.returncode not in (0, 3) or not unlocked.exists():
                raise RuntimeError(f"qpdf falhou: {r.stderr[:200]}")
            if c.dry:
                log("[dry] baixaria e salvaria", c.dest); pdf = unlocked; _ler_pdf(c, pdf); return
            c.dest.parent.mkdir(parents=True, exist_ok=True)
            c.dest.write_bytes(unlocked.read_bytes())
    _ler_pdf(c, c.dest)
    log("PDF ok:", c.dest.name)


def _ler_pdf(c, pdf):
    txt, secs = parse_pdf(pdf)
    venc = re.search(r"(\d\d)/(\d\d)/(\d{4})", txt[txt.index("Vencimento"):])
    if (int(venc[2]), int(venc[3])) != (c.mes, c.ano):
        raise RuntimeError(f"vencimento {venc[0]} não é de {c.ym} — PDF errado?")
    total = re.search(r"Total\s+R\$\s*([\d.]+,\d\d)", txt)
    linha = re.search(r"\d{5}\.\d{5} \d{5}\.\d{6} \d{5}\.\d{6} \d \d{6,14}", txt)
    d = c.d
    d["total"] = float(total[1].replace(".", "").replace(",", "."))
    d["venc_impresso"] = f"{venc[3]}-{venc[2]}-{venc[1]}"
    d["venc_pagto"] = dia_util(date(int(venc[3]), int(venc[2]), int(venc[1]))).isoformat()
    d["linha"] = (lambda s: s[:s.rindex(" ") + 1] + s.split()[-1].ljust(14, "0"))(linha[0]) if linha else None
    for k, (_, chave) in ABAS.items():
        sec = secao(secs, chave)
        soma = sum(t["valor"] for t in sec["tx"])
        if soma != sec["subtotal"]:
            raise RuntimeError(f"{k}: soma das transações {soma} ≠ subtotal do PDF {sec['subtotal']} — não escrevi nada")
        d[k.lower()] = float(soma)
    c._secs = secs


def saldo_xlsx(path):
    """Recalcula uma cópia no LibreOffice e lê Total!C13 (openpyxl não guarda valor de fórmula)."""
    import openpyxl
    with tempfile.TemporaryDirectory() as t:
        src = Path(t, "in.xlsx"); src.write_bytes(Path(path).read_bytes())
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", t + "/o", str(src)], capture_output=True, timeout=180)
        wb = openpyxl.load_workbook(Path(t, "o", "in.xlsx"), data_only=True)
        v = wb["Total"]["C13"].value
        out = Path(tempfile.gettempdir(), f"altus_recalc_{Path(path).name}")
        out.write_bytes(Path(t, "o", "in.xlsx").read_bytes())
    if v is None:
        raise RuntimeError("Total!C13 vazio após recalcular")
    return round(v, 2), out


def etapa_xlsx(c):
    import openpyxl
    from copy import copy
    from openpyxl.styles import PatternFill
    _ler_pdf(c, c.dest)
    path = c.xlsx()
    wb = openpyxl.load_workbook(path)
    escreveu = False
    for k, (aba, chave) in ABAS.items():
        if aba not in wb.sheetnames:
            raise RuntimeError(f"aba '{aba}' não existe em {path.name} (regra: o xlsx do mês nasce com as 3 abas)")
        ws = wb[aba]
        if ws["B5"].value is not None:
            continue  # já preenchida
        tpl = wb["BB Altus Visa Luiz"]
        r, primeiro = 5, True
        from itertools import groupby
        for cat, itens in groupby(secao(c._secs, chave)["tx"], key=lambda t: t["cat"]):
            if not primeiro:
                r += 2
            primeiro = False
            nome = cat or "Sem categoria"
            ws.cell(r, 2, nome[0]); ws.cell(r, 3, nome[1:]); r += 1
            for t in itens:
                for i, v in enumerate([t["data"], t["desc"], t["cidade"], t["pais"], float(t["valor"]), 0]):
                    ws.cell(r, 2 + i, v)
                r += 1
        ws["E2"] = f"=SUM(F5:F{max(300, r)})"
        escreveu = True
        log(f"aba {aba}: transações escritas até a linha {r - 1}")
    t = wb["Total"]
    for lin in range(5, 8):
        rot = str(t.cell(lin, 2).value or "")
        for k, (aba, _) in ABAS.items():
            if f"Altus {k}" in rot:
                t.cell(lin, 3).value = f"='{aba}'!E2"
                t.cell(lin, 3).fill = PatternFill(fill_type=None)
                escreveu = True
    if c.dry:
        c.d["saldo"] = saldo_xlsx(path)[0]
        log("[dry] escreveria no xlsx:", path.name, "| saldo atual do arquivo:", c.d["saldo"]); return
    if escreveu:
        wb.save(path)
    c.d["saldo"], _ = saldo_xlsx(path)
    log("xlsx ok; saldo (Total!C13) =", c.d["saldo"])


def etapa_boletos(c):
    import gspread
    pdf_link = drive_link(c.dest.name)
    gc = gspread.authorize(creds())
    sh = gc.open_by_key(BOLETOS_ID); ws = sh.worksheet(BOLETOS_TAB)
    rot = f"{MES[c.mes - 1].lower()}./{str(c.ano)[2:]}"
    lin = [i for i, r in enumerate(ws.get_all_values(), 1) if len(r) > 3 and r[1].strip() == rot and r[3].strip() == "BB Altus Visa"]
    if len(lin) != 1:
        raise RuntimeError(f"linha '{rot} / BB Altus Visa' na planilha Boletos: achei {len(lin)}")
    i = lin[0]
    if c.dry:
        log(f"[dry] Boletos linha {i}: dia {date.fromisoformat(c.d['venc_pagto']).day}, R$ {brl(c.d['total'])}, link {pdf_link}"); return
    ws.update(values=[[date.fromisoformat(c.d["venc_pagto"]).day]], range_name=f"C{i}", value_input_option="USER_ENTERED")
    ws.update(values=[[c.d["total"]]], range_name=f"E{i}", value_input_option="RAW")
    ws.update(values=[[pdf_link]], range_name=f"F{i}", value_input_option="USER_ENTERED")
    limpa_amarelo(sh, ws, [(i, 6)])
    log("Boletos linha", i, "atualizada")


def etapa_money(c):
    pdf_link = drive_link(c.dest.name)
    cmd = ["npx", "tsx", "--env-file=.env", "scripts/altus-conta-pagar.ts", str(c.d["total"]), c.d["venc_pagto"], pdf_link, str(c.dest.parent), c.d.get("linha") or ""]
    if c.dry:
        log("[dry] money:", cmd[3:]); return
    r = subprocess.run(cmd, cwd=MONEY, capture_output=True, text=True, timeout=180)
    if r.returncode:
        raise RuntimeError(f"money: {r.stderr[-300:]}")
    log("money:", r.stdout.strip()[-120:])


def etapa_pagtos(c):
    import gspread
    achados = sorted(c.pasta_mes.rglob("*Pagtos BB Altus Visa.gsheet"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not achados:
        raise Pendente("planilha 'Pagtos BB Altus Visa' do mês ainda não existe")
    doc = json.loads(achados[0].read_text())["doc_id"]
    pdf_link, xlsx_link = drive_link(c.dest.name), drive_link(c.xlsx().name)
    sh = gspread.authorize(creds()).open_by_key(doc); ws = sh.sheet1
    v = ws.get_all_values()
    nfc = lambda s: unicodedata.normalize("NFC", s).strip()
    hl = next((i for i, r in enumerate(v) if "Payee" in [x.strip() for x in r]), None)
    if hl is None:
        raise RuntimeError("planilha de pagtos: cabeçalho com 'Payee' não encontrado")
    col = {x.strip(): j + 1 for j, x in enumerate(v[hl]) if x.strip()}  # colunas por nome (a planilha ganha colunas)
    cv, cp, cl, cf = col["Valor"], col["Payee"], col["Link comprov"], col["Forma pagamento"]
    l_pais = [i for i, r in enumerate(v, 1) if len(r) >= cp and nfc(r[cp - 1]) == "Luiz e Fátima"]
    l_tot = [i for i, r in enumerate(v, 1) if len(r) > 1 and "Total da fatura" in r[1]]
    if len(l_pais) != 1 or len(l_tot) != 1:
        raise RuntimeError("planilha de pagtos: não achei 1 linha 'Luiz e Fátima' + 1 'Total da fatura'")
    a, b = l_pais[0], l_tot[0]
    A = lambda r, cc: gspread.utils.rowcol_to_a1(r, cc)
    if c.dry:
        log(f"[dry] pagtos: linha {a} R$ {brl(c.d['luiz'] + c.d['fátima'])} (col {cv}) + link (col {cl}); linha {b} total R$ {brl(c.d['total'])} + pdf (col {cf})"); return
    ws.update(values=[[round(c.d["luiz"] + c.d["fátima"], 2)]], range_name=A(a, cv), value_input_option="RAW")
    ws.update(values=[[xlsx_link]], range_name=A(a, cl), value_input_option="USER_ENTERED")
    ws.update(values=[[c.d["total"]]], range_name=A(b, cv), value_input_option="RAW")
    ws.update(values=[[pdf_link]], range_name=A(b, cf), value_input_option="USER_ENTERED")
    limpa_amarelo(sh, ws, [(a, cv), (a, cl), (b, cv), (b, cf)])
    log("pagtos atualizada")


# ---- templates (e-mail / WhatsApp) — texto vive nos .txt da pasta Templates (Fábio edita lá)
def _template(nome):
    return (TEMPLATES / nome).read_text()


def _preenche(s, c):
    m_ref = MES[(c.mes - 2) % 12]
    venc = date.fromisoformat(c.d["venc_pagto"])
    vals = {"MES_ENVIO": MES[c.mes - 1], "MES_REF": m_ref, "ANO2": str(c.ano)[2:], "VALOR_A_PAGAR": brl(c.d["saldo"]),
            "DATA_PAGTO": f"{venc.day} {MES[venc.month - 1]}", "LINHA_DIGITAVEL_BOLETO": c.d["linha"]}
    for k, v in vals.items():
        s = s.replace("{" + k + "}", v)
    return s


def _confere(*partes):
    if any(re.search(r"\{[A-Z_0-9]+\}", p) for p in partes):
        raise RuntimeError("placeholder sem preencher no template")


def _secao(t, ini, fim="--- VARIÁVEIS"):
    a = t.index(ini); a = t.index("\n", a) + 1
    return t[a:t.index(fim)].strip("\n")


def etapa_email(c):
    t = _preenche(_template("Template - Email fatura cartões aos pais (BTG e BB).txt"), c)
    head = t[:t.index("--- CORPO")]
    para = re.search(r"PARA:(.*?)ASSUNTO:", head, re.S)[1]
    destinos = re.findall(r"[\w.+-]+@[\w.-]+\.\w+", para)
    assunto = re.search(r"ASSUNTO:\s*(.+?)\s{2,}\(", head)[1]
    corpo = _secao(t, "--- CORPO")
    _confere(assunto, corpo)
    corpo = corpo.replace("[assinatura HTML padrão do Fábio]", "").replace("> Abs,", "Abs,").rstrip()
    from html import escape
    h = "".join("<p style=\"margin:0 0 14px\">" + re.sub(r"\*(.+?)\*", r"<b><u>\1</u></b>", escape(p).replace("\n", "<br>")) + "</p>" for p in corpo.split("\n\n"))
    sig = re.search(r"ASSINATURA_FABIO_FORMAL = `(.*?)`;", (KPOOL / "src/lib/email.ts").read_text(), re.S)[1]
    html = f'<div style="font-family:Arial,Helvetica,sans-serif;font-size:14px;color:#222">{h}{sig}</div>'
    _, anexo = saldo_xlsx(c.xlsx())  # cópia recalculada: prévia no celular mostra os valores
    if c.dry:
        log("[dry] e-mail:", destinos, "|", assunto, "| anexo", c.xlsx().name); print(corpo); return
    env = dotenv_values(KPOOL / ".env.local")
    body = {"to": destinos, "from": "Fábio Póvoa <fabio@smartmoney.ventures>", "subject": assunto, "corpoHtml": html,
            "descricao": f"Fatura cartões {c.ym}", "attachmentBase64": base64.b64encode(anexo.read_bytes()).decode(), "attachmentName": c.xlsx().name}
    req = urllib.request.Request("https://kpool.smartmoney.ventures/api/cron/enviar-email", json.dumps(body).encode(),
                                 {"Authorization": f"Bearer {env['CRON_SECRET']}", "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=90) as r:
        log("e-mail enviado:", r.status, assunto)


def etapa_whatsapp(c):
    t = _preenche(_template("Template - WhatsApp fatura cartões aos pais (BTG e BB).txt"), c)
    jid = re.search(r"JID: (\d+@g\.us)", t)[1]
    msg = _secao(t, "--- MENSAGEM")
    _confere(msg)
    if c.dry:
        log("[dry] WhatsApp →", jid); print(msg); return
    env = dotenv_values(KPOOL / ".env.local")
    req = urllib.request.Request(f"{env['EVOLUTION_URL']}/message/sendText/fabio", json.dumps({"number": jid, "text": msg}).encode(),
                                 {"apikey": env["EVOLUTION_API_KEY"], "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as r:
        log("WhatsApp enviado:", r.status, json.load(r).get("key", {}).get("id"))


ETAPAS = [("pdf", etapa_pdf), ("xlsx", etapa_xlsx), ("boletos", etapa_boletos), ("money", etapa_money),
          ("pagtos", etapa_pagtos), ("email", etapa_email), ("whatsapp", etapa_whatsapp)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--month"); ap.add_argument("--dry-run", action="store_true"); ap.add_argument("--mark-done", action="store_true")
    a = ap.parse_args()
    hoje = date.today()
    ano, mes = map(int, a.month.split("-")) if a.month else (hoje.year, hoje.month)
    c = Ctx(ano, mes, a.dry_run)
    if a.mark_done:
        c.st["steps"] = {n: "marcado manualmente" for n, _ in ETAPAS}; c.save(); log("marcado como concluído:", c.ym); return
    pend = [n for n, _ in ETAPAS if n not in c.st["steps"]]
    if not pend and not a.dry_run:
        return
    log("mês", c.ym, "pendentes:", pend, "(dry-run)" if a.dry_run else "")
    for nome, fn in ETAPAS:
        if nome in c.st["steps"] and not a.dry_run:
            continue
        if nome not in ("pdf", "xlsx") and not (c.st["steps"].get("xlsx") or a.dry_run):
            break
        try:
            fn(c)
            if not a.dry_run:
                c.st["steps"][nome] = datetime.now().isoformat(timespec="seconds"); c.save()
        except Pendente as e:
            log(f"[{nome}] pendente: {e}")
            if nome in ("pdf", "xlsx"):
                if nome == "pdf" and hoje.day == 8 and not a.dry_run:
                    notify("Fatura Altus ainda não chegou por e-mail (dia 8).")
                break
        except Exception as e:  # noqa: BLE001
            log(f"[{nome}] ERRO: {e}")
            notify(f"Altus {c.ym}: erro na etapa {nome} — {e}")
            if nome in ("pdf", "xlsx"):
                break
    if not a.dry_run and all(n in c.st["steps"] for n, _ in ETAPAS):
        notify(f"Fatura Altus {c.ym} processada: PDF, xlsx, Boletos, money, pagtos, e-mail e WhatsApp.")


if __name__ == "__main__":
    main()
